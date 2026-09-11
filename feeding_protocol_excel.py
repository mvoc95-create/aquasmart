"""Excel interchange for the existing feeding table (no database writes here)."""
import hashlib
import io
import json
import math
import unicodedata
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
FEED_PREFIX = 'Ração (g): '
COLUMNS = [
    ('ID', 'id'), ('Ativa', 'active'), ('Fase real', 'phase'),
    ('Dia fase', 'phase_day'), ('Dia ciclo', 'cycle_day'), ('Estágio', 'stage_label'),
    ('População', 'population'), ('Sobrev. %', 'survival_pct'),
    ('Peso tabela (g)', 'individual_weight_g'), ('Taxa alim. %', 'feed_rate_pct'),
    ('Total/dia (g)', 'total_day_g'), ('Tratos/dia', 'feedings_per_day'),
]
PHASE_NAMES = {'bercario': 'Berçário', 'juvenil': 'Juvenil', 'engorda': 'Engorda'}


def normalized(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value).strip().lower())
                   if not unicodedata.combining(c))


def snapshot(rows):
    """Detect edits between preview and confirmation, including a PDF reset."""
    return hashlib.sha256(json.dumps(rows, sort_keys=True, ensure_ascii=False,
                                     default=str).encode()).hexdigest()


def export_table(rows, labels):
    # Runtime export follows the application's existing openpyxl dependency.
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tabela base'
    ws.append([title for title, _ in COLUMNS] + [FEED_PREFIX + label for label in labels])
    for row in rows:
        values = [row[key] for _, key in COLUMNS]
        values[1] = 'Sim' if row['active'] else 'Não'
        values[2] = PHASE_NAMES.get(row['phase'], row['phase'])
        mix = {item['label']: item['grams'] for item in row['mixes']}
        ws.append(values + [mix.get(label, 0) for label in labels])
        # Labels are text, even if a user-entered label starts with '='.
        for col in (2, 3, 6):
            ws.cell(ws.max_row, col).data_type = 's'
    ws.freeze_panes = 'G2'
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 44
    for cell in ws[1]:
        cell.data_type = 's'
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176B70')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for index in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(index)].width = 23 if index > 12 else 18
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.font = Font(color='808080' if cell.column == 1 else '154F91')
            if cell.column in (8, 9, 10):
                cell.number_format = '0.0000'
    notes = wb.create_sheet('Como importar')
    instructions = [
        'Aqua Smart — atualização da tabela base',
        'Edite a aba Tabela base e mantenha os cabeçalhos e o ID de cada linha.',
        'O ID identifica a linha mesmo depois de ordenar a planilha. Não copie IDs entre linhas.',
        'Você pode manter somente as linhas que deseja atualizar. Linhas ausentes não serão excluídas.',
        'Informe gramas nas colunas de ração. Célula vazia de ração significa zero.',
        'Mantenha todas as colunas de ração existentes. Para nova ração, acrescente Ração (g): NOME.',
        'Percentuais: digite 5 para 5%. Também são aceitas células formatadas como 5%.',
        'Use valores, não fórmulas. No Excel, copie e cole como Valores antes de importar.',
        'Ativa: Sim ou Não. Fase real: Berçário, Juvenil ou Engorda.',
        'Engorda mantém os 4 tratos diários do sistema. Nas demais fases, use de 1 a 24.',
        'Biomassa é recalculada. O total diário só é recalculado se marcar essa opção na importação.',
        'Os valores de cada ração definem o mix, como na edição manual. Não são redistribuídos automaticamente.',
        'A importação não altera vínculos com estoque, ajustes por unidade ou configurações de biometria.',
        'Novas rações ficam sem vínculo: associe ao estoque no cabeçalho da tabela após importar.',
        'Baixe uma cópia atual antes de editar. Revise a prévia e confirme para gravar.',
    ]
    for note in instructions:
        notes.append([note])
    notes.column_dimensions['A'].width = 110
    for cells in notes:
        cells[0].alignment = Alignment(wrap_text=True, vertical='center')
        notes.row_dimensions[cells[0].row].height = 32
    notes['A1'].font = Font(bold=True, size=16, color='176B70')
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def number(cell, title, minimum=0, maximum=2147483647, integer=False, percent=False, blank_zero=False):
    value = cell.value
    if value is None or str(value).strip() == '':
        if blank_zero:
            return 0
        raise ValueError(f'{title}: valor obrigatório.')
    if isinstance(value, bool):
        raise ValueError(f'{title}: informe um número.')
    text = str(value).strip()
    explicit_pct = text.endswith('%')
    if explicit_pct:
        if not percent:
            raise ValueError(f'{title}: não aceita percentual.')
        text = text[:-1].strip()
    if ',' in text:
        text = text.replace('.', '').replace(',', '.')
    try:
        result = float(text)
    except (ValueError, TypeError):
        raise ValueError(f'{title}: número inválido.') from None
    if percent and not explicit_pct and '%' in cell.number_format:
        result *= 100
    if not math.isfinite(result) or not minimum <= result <= maximum:
        raise ValueError(f'{title}: informe um valor entre {minimum} e {maximum}.')
    if integer and not result.is_integer():
        raise ValueError(f'{title}: informe um número inteiro.')
    return int(result) if integer else result


def parse_table(data, filename, current_rows, labels, recalc=False):
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('Selecione um arquivo Excel .xlsx. Salve arquivos .xls como .xlsx no Excel.')
    if not data or len(data) > MAX_BYTES:
        raise ValueError('O arquivo deve ter até 5 MB e não pode estar vazio.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('Planilha muito grande. Use o modelo baixado pelo sistema.')
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception:
        raise ValueError('Não foi possível ler o Excel. Envie uma planilha .xlsx válida.') from None
    try:
        if 'Tabela base' not in wb.sheetnames:
            raise ValueError('A aba Tabela base não foi encontrada. Use o modelo baixado pelo sistema.')
        ws = wb['Tabela base']
        if ws.max_row > MAX_ROWS + 1 or ws.max_column > 150:
            raise ValueError('Limite: 2.000 linhas e 150 colunas. Remova linhas/colunas extras do Excel.')
        iterator = ws.iter_rows()
        header = [str(c.value).strip() if c.value is not None else '' for c in next(iterator)]
        while header and not header[-1]:
            header.pop()
        if len(header) != len(set(header)) or '' in header:
            raise ValueError('Há cabeçalhos vazios ou repetidos na primeira linha.')
        required = [title for title, _ in COLUMNS] + [FEED_PREFIX + label for label in labels]
        missing = [title for title in required if title not in header]
        if missing:
            raise ValueError('Colunas ausentes: ' + ', '.join(missing) + '. Baixe o modelo atualizado.')
        feed_labels = []
        for title in header:
            if title in [name for name, _ in COLUMNS]:
                continue
            if not title.startswith(FEED_PREFIX) or not title[len(FEED_PREFIX):].strip():
                raise ValueError(f'Coluna desconhecida: {title}. Use o modelo do sistema.')
            label = title[len(FEED_PREFIX):]
            if len(label) > 160 or label != label.strip():
                raise ValueError('Nome de ração inválido (máximo de 160 caracteres, sem espaços nas pontas).')
            feed_labels.append(label)
        if len({normalized(label) for label in feed_labels}) != len(feed_labels):
            raise ValueError('Há nomes de rações repetidos ou equivalentes.')
        existing = {row['id']: row for row in current_rows}
        seen, parsed, errors = set(), [], []
        for line_number, cells in enumerate(iterator, start=2):
            if all(c.value is None or str(c.value).strip() == '' for c in cells):
                continue
            try:
                if any(c.data_type in ('f', 'e') for c in cells):
                    raise ValueError('Fórmula ou erro do Excel: copie e cole como Valores antes de importar.')
                row = dict(zip(header, cells))
                rid = number(row['ID'], 'ID', minimum=1, integer=True)
                if rid not in existing:
                    raise ValueError(f'ID {rid} não existe. Baixe o modelo atual; esta importação atualiza linhas existentes.')
                if rid in seen:
                    raise ValueError(f'ID {rid} repetido.')
                seen.add(rid)
                active = normalized(row['Ativa'].value)
                if active not in ('sim', 'nao', '1', '0', 'true', 'false'):
                    raise ValueError('Ativa: informe Sim ou Não.')
                phase = normalized(row['Fase real'].value)
                if phase not in PHASE_NAMES:
                    raise ValueError('Fase real: use Berçário, Juvenil ou Engorda.')
                stage = str(row['Estágio'].value or '').strip()
                if not stage or len(stage) > 50:
                    raise ValueError('Estágio: informe de 1 a 50 caracteres.')
                result = {'id': rid, 'active': active in ('sim', '1', 'true'),
                          'phase': phase, 'stage_label': stage}
                for title, key in COLUMNS:
                    if key in result:
                        continue
                    result[key] = number(
                        row[title], title, minimum=1 if key in ('phase_day', 'cycle_day', 'feedings_per_day') else 0,
                        maximum=100 if key in ('survival_pct', 'feed_rate_pct') else (24 if key == 'feedings_per_day' else 2147483647),
                        integer=key in ('phase_day', 'cycle_day', 'population', 'total_day_g', 'feedings_per_day'),
                        percent=key in ('survival_pct', 'feed_rate_pct'))
                if phase == 'engorda' and result['feedings_per_day'] != 4:
                    raise ValueError('A engorda utiliza 4 tratos/dia. Informe 4.')
                result['biomass_kg'] = round(result['population'] * result['individual_weight_g'] / 1000, 3)
                if recalc:
                    result['total_day_g'] = round(result['biomass_kg'] * result['feed_rate_pct'] * 10)
                    if result['total_day_g'] > 2147483647:
                        raise ValueError('Total diário recalculado acima do limite permitido.')
                result['mixes'] = [
                    {'label': label, 'grams': number(row[FEED_PREFIX + label], label, integer=True, blank_zero=True)}
                    for label in feed_labels
                ]
                parsed.append(result)
            except ValueError as error:
                errors.append(f'Linha {line_number}: {error}')
                if len(errors) >= 15:
                    break
        if errors:
            raise ValueError('\n'.join(errors))
        if not parsed:
            raise ValueError('A planilha não contém linhas de dados.')
        return parsed, feed_labels
    finally:
        wb.close()
